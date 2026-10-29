"""用完整保存重开的PPTX核文本、占位符、合并和图表交付。"""
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pptx
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def saved(presentation) -> BytesIO:
    """保存合成演示包。@param presentation 实际演示文稿。@return 可重开的内存流。"""
    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    return output


def main() -> None:
    """逐项核候选源码及OOXML，不启动Office。@return None。"""
    root = Path(os.environ['PROJECT_CANDIDATE_ROOT']).resolve()
    result = {'candidate_import': Path(pptx.__file__).resolve() == root / 'library/src/pptx/__init__.py'}
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    task = sys.argv[1]
    if task == 'PP-U01':
        first = slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text_frame
        first.text = 'A\nB\vC'
        second = slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text_frame
        second.paragraphs[0].text = 'A\nB\vC'
        output = saved(presentation)
        frames = [shape.text_frame for shape in Presentation(output).slides[0].shapes]
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        with ZipFile(output) as archive:
            xml = ET.fromstring(archive.read('ppt/slides/slide1.xml'))
        result['text_frame_newlines_create_paragraphs_while_paragraph_newlines_create_breaks'] = (
            len(frames[0].paragraphs) == 2 and frames[0].text == 'A\nB\vC'
            and len(frames[1].paragraphs) == 1 and frames[1].text == 'A\vB\vC'
            and len(xml.findall('.//a:br', ns)) == 3)
    elif task == 'PP-C01':
        first = slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text_frame
        old = first.paragraphs[0].add_run()
        old.text, old.font.bold = 'OLD', True
        first.text = 'REPLACED'
        second = slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text_frame
        run = second.paragraphs[0].add_run()
        run.text, run.font.bold = 'OLD', True
        run.text = 'NEW'
        reopened = Presentation(saved(presentation)).slides[0]
        result['frame_replacement_discards_direct_run_formatting_while_run_replacement_preserves_it'] = (
            reopened.shapes[0].text_frame.paragraphs[0].runs[0].font.bold is None
            and reopened.shapes[1].text_frame.paragraphs[0].runs[0].font.bold is True
            and reopened.shapes[1].text == 'NEW')
    elif task == 'PP-U02':
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        # 仅构造索引有空洞的合成模板，不改候选源码或原版模板文件。
        slide.placeholders[1]._element.xpath('.//p:ph')[0].set('idx', '7')
        placeholders = Presentation(saved(presentation)).slides[0].placeholders
        missing = False
        try:
            placeholders[1]
        except KeyError:
            missing = True
        result['placeholder_lookup_uses_template_idx_not_collection_position'] = (
            len(placeholders) == 2 and [item.placeholder_format.idx for item in placeholders] == [0, 7]
            and placeholders[7].placeholder_format.idx == 7 and missing)
    elif task == 'PP-C02':
        table = slide.shapes.add_table(1, 2, 0, 0, Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = 'A', 'B'
        table.cell(0, 0).merge(table.cell(0, 1))
        output = saved(presentation)
        table = Presentation(output).slides[0].shapes[0].table
        cells = list(table.iter_cells())
        before = (cells[0].text, cells[1].text, cells[0].span_width)
        distinct = cells[0]._tc is not cells[1]._tc
        flags = cells[0].is_merge_origin and not cells[0].is_spanned and cells[1].is_spanned
        cells[0].split()
        result['merged_pptx_cells_keep_grid_slots_and_split_does_not_restore_moved_text'] = (
            len(cells) == 2 and distinct and flags and before == ('A\nB', '', 2)
            and table.cell(0, 0).text == 'A\nB' and table.cell(0, 1).text == ''
            and not table.cell(0, 0).is_merge_origin and not table.cell(0, 1).is_spanned)
    elif task == 'PP-P01':
        original, replacement = CategoryChartData(), CategoryChartData()
        original.categories = ['Q1', 'Q2']
        original.add_series('Revenue', [1, 2])
        replacement.categories = ['Q1', 'Q2']
        replacement.add_series('Revenue', [30, 40])
        graphic = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), original)
        graphic.chart.replace_data(replacement)
        output = saved(presentation)
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
            's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with ZipFile(output) as archive:
            chart = ET.fromstring(archive.read('ppt/charts/chart1.xml'))
            embedded = [name for name in archive.namelist() if name.startswith('ppt/embeddings/') and name.endswith('.xlsx')]
            with ZipFile(BytesIO(archive.read(embedded[0]))) as workbook:
                worksheet = ET.fromstring(workbook.read('xl/worksheets/sheet1.xml'))
        cache = [item.text for item in chart.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v', ns)]
        sheet_values = {cell.get('r'): cell.findtext('s:v', namespaces=ns) for cell in worksheet.findall('.//s:c', ns)}
        result['replace_data_updates_chart_cache_and_embedded_workbook_without_office_recalculation'] = (
            cache == ['30', '40'] and sheet_values.get('B2') == '30' and sheet_values.get('B3') == '40'
            and len(embedded) == 1 and Presentation(output).slides[0].shapes[0].has_chart)
    else:
        raise ValueError('未冻结的业务任务')
    print(json.dumps(result))


if __name__ == '__main__':
    main()

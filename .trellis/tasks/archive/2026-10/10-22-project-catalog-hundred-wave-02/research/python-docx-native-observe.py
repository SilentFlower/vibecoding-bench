"""以保存重开及独立DOCX部件核文本、表格、日期与批注归属。"""
import datetime as dt
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import docx
from docx import Document


NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'd': 'http://purl.org/dc/terms/'}


def saved(document) -> BytesIO:
    """保存合成文档，不写真实模板。@param document 文档。@return 可重开的内存包。"""
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def main() -> None:
    """执行单条固定候选观察，不启动Word或LibreOffice。@return None。"""
    root = Path(os.environ['PROJECT_CANDIDATE_ROOT']).resolve()
    result = {'candidate_import': Path(docx.__file__).resolve() == root / 'library/src/docx/__init__.py'}
    task = sys.argv[1]
    if task == 'DX-U01':
        document = Document()
        paragraph = document.add_paragraph(style='Heading 1')
        paragraph.add_run('A').bold = True
        paragraph.add_run('B').italic = True
        paragraph.text = 'REPLACED'
        control = document.add_paragraph()
        run = control.add_run('OLD')
        run.bold = True
        run.text = 'NEW'
        output = saved(document)
        reopened = Document(output)
        with ZipFile(output) as archive:
            xml = ET.fromstring(archive.read('word/document.xml'))
        first = xml.find('w:body/w:p', NS)
        result['paragraph_text_replaces_runs_while_run_text_preserves_its_formatting'] = (
            len(reopened.paragraphs[0].runs) == 1 and reopened.paragraphs[0].runs[0].bold is None
            and reopened.paragraphs[0].runs[0].italic is None and first.find('w:pPr/w:pStyle', NS) is not None
            and reopened.paragraphs[1].runs[0].text == 'NEW' and reopened.paragraphs[1].runs[0].bold is True)
    elif task == 'DX-U02':
        document = Document()
        document.add_paragraph('A')
        document.add_table(1, 1).cell(0, 0).text = 'T'
        document.add_paragraph('B')
        reopened = Document(saved(document))
        items = list(reopened.iter_inner_content())
        result['ordered_block_iteration_preserves_paragraph_table_interleaving'] = (
            [type(item).__name__ for item in items] == ['Paragraph', 'Table', 'Paragraph']
            and [p.text for p in reopened.paragraphs] == ['A', 'B'] and len(reopened.tables) == 1
            and items[1].cell(0, 0).text == 'T')
    elif task == 'DX-C01':
        document = Document()
        table = document.add_table(1, 2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'A', 'B'
        table.cell(0, 0).merge(table.cell(0, 1))
        output = saved(document)
        reopened = Document(output)
        cells = reopened.tables[0].rows[0].cells
        with ZipFile(output) as archive:
            row = ET.fromstring(archive.read('word/document.xml')).find('.//w:tr', NS)
        result['merged_cells_repeat_layout_positions_but_have_one_underlying_xml_cell'] = (
            len(cells) == 2 and cells[0]._tc is cells[1]._tc and cells[0].grid_span == 2
            and cells[0].text == cells[1].text == 'A\nB' and len(row.findall('w:tc', NS)) == 1)
    elif task == 'DX-C02':
        aware = dt.datetime(2026, 4, 3, 10, tzinfo=dt.timezone(dt.timedelta(hours=8)))
        raw, normalized = Document(), Document()
        raw.core_properties.created = aware
        normalized.core_properties.created = aware.astimezone(dt.timezone.utc)
        raw_buffer, normalized_buffer = saved(raw), saved(normalized)
        with ZipFile(raw_buffer) as archive:
            raw_xml = ET.fromstring(archive.read('docProps/core.xml')).findtext('d:created', namespaces=NS)
        result['aware_core_datetime_requires_explicit_utc_normalization_at_the_caller'] = (
            raw_xml == '2026-04-03T10:00:00Z'
            and Document(raw_buffer).core_properties.created == dt.datetime(2026, 4, 3, 10, tzinfo=dt.timezone.utc)
            and Document(normalized_buffer).core_properties.created == aware.astimezone(dt.timezone.utc))
    elif task == 'DX-P01':
        first, second = Document(), Document()
        foreign = first.add_paragraph('A').runs[0]
        second.add_comment(foreign, 'LOCAL_REVIEW', author='Synthetic')
        first_buffer, second_buffer = saved(first), saved(second)
        with ZipFile(first_buffer) as archive:
            first_xml = ET.fromstring(archive.read('word/document.xml'))
            dangling = first_xml.find('.//w:commentReference', NS) is not None and 'word/comments.xml' not in archive.namelist()
        with ZipFile(second_buffer) as archive:
            second_has_part = 'word/comments.xml' in archive.namelist()
        control = Document()
        local = control.add_paragraph('A').runs[0]
        owned = local.part is control.part and foreign.part is not second.part
        control.add_comment(local, 'LOCAL_REVIEW', author='Synthetic')
        with ZipFile(saved(control)) as archive:
            control_xml = ET.fromstring(archive.read('word/document.xml'))
            valid = 'word/comments.xml' in archive.namelist() and control_xml.find('.//w:commentReference', NS) is not None
        result['comment_runs_must_belong_to_the_target_document_before_publication'] = dangling and second_has_part and owned and valid
    else:
        raise ValueError('未冻结的业务任务')
    print(json.dumps(result))


if __name__ == '__main__':
    main()

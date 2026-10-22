"""核表格导入、输出格式器、批量追加及按字段名对齐修复。"""
from json import dumps
from os import environ
from pathlib import Path
from sys import argv

import tablib
from tablib import Dataset
from tablib.exceptions import InvalidDimensions, UnsupportedFormat


def main() -> None:
    """从实际候选核关键数据与拒绝路径，不读取参考答案。@return None。"""
    root = Path(environ['PROJECT_CANDIDATE_ROOT']).resolve()
    result = {'candidate_import': Path(tablib.__file__).resolve() == root / 'library/src/tablib/__init__.py'}
    task = argv[1]
    if task == 'TB-U01':
        data = Dataset().load('sku,qty\nA,2\nB,\n', format='csv')
        result['csv_headers_and_cell_values_are_explicit_text_contracts'] = (
            data.headers == ['sku', 'qty'] and data.dict == [{'sku': 'A', 'qty': '2'}, {'sku': 'B', 'qty': ''}]
            and Dataset().load('sku,qty\nA,2\n', format='csv', headers=False)[0] == ('sku', 'qty'))
    elif task == 'TB-U02':
        data = Dataset(('A', 2), headers=['sku', 'qty'])
        data.add_formatter('qty', lambda value: f'{value:03}')
        result['export_formatters_transform_projection_without_changing_stored_cells'] = (
            data[0] == ('A', 2) and data.dict == [{'sku': 'A', 'qty': '002'}]
            and data.export('csv') == 'sku,qty\r\nA,002\r\n' and data[0] == ('A', 2))
    elif task == 'TB-C01':
        from io import BytesIO
        from openpyxl import load_workbook

        data = Dataset(('=1+1',), headers=['note'])
        ordinary = load_workbook(BytesIO(data.export('xlsx'))).active['A2']
        escaped = load_workbook(BytesIO(data.export('xlsx', escape=True))).active['A2']
        result['xlsx_escape_is_an_explicit_lossy_formula_export_policy'] = (
            ordinary.value == '=1+1' and ordinary.data_type == 'f'
            and escaped.value == '1+1' and escaped.data_type == 's'
            and data[0] == ('=1+1',))
    elif task == 'TB-P01':
        data = Dataset(('A', 2), headers=['sku', 'qty'])
        raised = False
        try:
            data.extend([('B', 3), ('C',)])
        except InvalidDimensions:
            raised = True
        result['extend_can_keep_a_successful_prefix_after_a_later_invalid_row'] = (
            raised and data.dict == [{'sku': 'A', 'qty': 2}, {'sku': 'B', 'qty': 3}])
    elif task == 'TB-M01':
        source = [{'sku': 'A', 'qty': 2}, {'qty': 5, 'sku': 'B'}]
        original = [dict(row) for row in source]
        direct = Dataset()
        direct.dict = source
        expected = [{'sku': 'A', 'qty': 2}, {'sku': 'B', 'qty': 5}]
        result['mapping_rows_follow_header_names_without_mutating_input'] = direct.dict == expected and source == original
        parsed = Dataset().load('[{"sku":"A","qty":2},{"qty":5,"sku":"B"}]', format='json')
        yaml = Dataset().load('- sku: A\n  qty: 2\n- qty: 5\n  sku: B\n', format='yaml')
        result['json_and_yaml_share_the_repaired_mapping_alignment'] = parsed.dict == yaml.dict == expected
        failures = []
        for rows in ([{'sku': 'A', 'qty': 2}, {'sku': 'B'}],
                [{'sku': 'A', 'qty': 2}, {'sku': 'B', 'qty': 3, 'extra': 4}],
                [{'sku': 'A', 'qty': 2}, ['B', 3]]):
            target = Dataset(('old', 9), headers=['sku', 'qty'])
            failed = False
            try:
                target.dict = rows
            except UnsupportedFormat:
                failed = True
            failures.append(failed and target.dict == [{'sku': 'old', 'qty': 9}])
        unchanged = Dataset(('old', 9), headers=['sku', 'qty'])
        unchanged.dict = []
        lists = Dataset()
        lists.dict = [['A', 2], ['B', 5]]
        result['invalid_mapping_schemas_are_rejected_before_replacing_existing_data'] = all(failures)
        result['empty_list_and_positional_list_compatibility_are_preserved'] = (
            unchanged.dict == [{'sku': 'old', 'qty': 9}] and list(lists) == [('A', 2), ('B', 5)])
    else:
        raise ValueError('未冻结的业务任务')
    print(dumps(result))


if __name__ == '__main__':
    main()

"""图片说明和标题使用真实wp:docPr，无需视觉猜测。"""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from docx import Document

def test_inline_picture_accessibility_survives_docx_roundtrip():
    """保图片关系、尺寸和未知属性，仅编辑对应说明。@return None。"""
    document=Document();picture=document.add_picture(str(Path(__file__).resolve().parent/'test_files'/'python-icon.png'))
    width,height=picture.width,picture.height
    picture.alternative_text='中文说明 & details';picture.title='Figure 1'
    output=BytesIO();document.save(output);output.seek(0);restored=Document(output).inline_shapes[0]
    assert restored.alternative_text=='中文说明 & details' and restored.title=='Figure 1'
    assert restored.width==width and restored.height==height

def test_empty_is_explicit_none_removes_and_other_metadata_stays():
    """未设置与故意空说明可区分，不改绘图id/name。@return None。"""
    document=Document();picture=document.add_picture(str(Path(__file__).resolve().parent/'test_files'/'python-icon.png'))
    original=picture._inline.docPr.get('name')
    assert picture.alternative_text is None
    picture.alternative_text='';picture.title='keep'
    assert picture.alternative_text=='';picture.alternative_text=None
    assert picture.alternative_text is None and picture.title=='keep' and picture._inline.docPr.get('name')==original


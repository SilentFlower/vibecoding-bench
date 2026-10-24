"Public regression tests for PYTHON-DOCX-02-DEV; the complete business contract is in materials/PYTHON-DOCX-02-DEV-contract.md."
from datetime import datetime,timezone,timedelta
from io import BytesIO
from zipfile import ZipFile
from docx import Document

def test_offset_utc_conversion_is_visible_in_real_core_xml():
    "Verify offset utc conversion is visible in real core xml."
    document=Document();value=datetime(2024,1,2,1,30,15,123456,tzinfo=timezone(timedelta(hours=5,minutes=30)))
    document.core_properties.created=value;output=BytesIO();document.save(output)
    with ZipFile(output) as archive:assert b'2024-01-01T20:00:15Z' in archive.read('docProps/core.xml')
    output.seek(0);restored=Document(output).core_properties.created
    assert restored.replace(tzinfo=timezone.utc)==value.astimezone(timezone.utc).replace(microsecond=0)

def test_negative_offset_naive_and_all_properties_keep_precision_policy():
    "Verify negative offset naive and all properties keep precision policy."
    document=Document()
    for name in ('created','modified','last_printed'):
        value=datetime(2024,2,29,23,59,3,tzinfo=timezone(timedelta(hours=-4)))
        setattr(document.core_properties,name,value)
        assert getattr(document.core_properties,name)==datetime(2024,3,1,3,59,3,tzinfo=timezone.utc)
    document.core_properties.created=datetime(2024,1,1,8,0,0)
    assert document.core_properties.created==datetime(2024,1,1,8,tzinfo=timezone.utc)


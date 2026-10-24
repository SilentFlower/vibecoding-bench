"Public regression tests for PYTHON-DOCX-05-DEV; the complete business contract is in materials/PYTHON-DOCX-05-DEV-contract.md."
from io import BytesIO
import unittest
from docx import Document

def test_two_header_rows_data_cannot_split_and_real_roundtrip():
    "Verify two header rows data cannot split and real roundtrip."
    document=Document();table=document.add_table(rows=3,cols=2)
    table.rows[0].repeat_header=True;table.rows[1].repeat_header=True;table.rows[2].allow_split_across_pages=False
    output=BytesIO();document.save(output);output.seek(0);restored=Document(output).tables[0]
    assert restored.rows[0].repeat_header is True and restored.rows[1].repeat_header is True
    assert restored.rows[2].allow_split_across_pages is False
    assert 'cantSplit' in restored.rows[2]._tr.xml

def test_false_and_none_clear_inheritance_without_touching_height():
    "Verify false and none clear inheritance without touching height."
    row=Document().add_table(rows=1,cols=1).rows[0]
    assert row.repeat_header is None and row.allow_split_across_pages is None
    row.repeat_header=False;row.allow_split_across_pages=True
    assert row.repeat_header is False and row.allow_split_across_pages is True
    row.repeat_header=None;row.allow_split_across_pages=None
    assert row.repeat_header is None and row.allow_split_across_pages is None
    with unittest.TestCase().assertRaises(TypeError):row.repeat_header=1


"Public regression tests for PYTHON-DOCX-01-DEV; the complete business contract is in materials/PYTHON-DOCX-01-DEV-contract.md."
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement

def test_empty_created_modified_last_printed_roundtrip():
    "Verify empty created modified last printed roundtrip."
    document=Document();core=document.core_properties._element
    for name in ('created','modified','lastPrinted'):
        element=core._get_or_add(name);element.text=None
    output=BytesIO();document.save(output);output.seek(0);restored=Document(output)
    assert restored.core_properties.created is None and restored.core_properties.modified is None and restored.core_properties.last_printed is None

def test_whitespace_invalid_missing_and_legal_date_keep_policy():
    "Verify whitespace invalid missing and legal date keep policy."
    document=Document();core=document.core_properties._element
    core._get_or_add('created').text='   ';assert document.core_properties.created is None
    core._get_or_add('modified').text='not-a-date';assert document.core_properties.modified is None
    core._get_or_add('created').text='2024-02-29';assert document.core_properties.created.year==2024


"Public regression tests for PYTHON-DOCX-04-DEV; the complete business contract is in materials/PYTHON-DOCX-04-DEV-contract.md."
from io import BytesIO
from zipfile import ZipFile
from docx import Document

def test_hyperlink_roundtrip_external_internal_text_and_run_format():
    "Verify hyperlink roundtrip external internal text and run format."
    document=Document();paragraph=document.add_paragraph()
    link=paragraph.add_hyperlink('Original','https://example.invalid/path',fragment='section')
    link.runs[0].bold=True;link.text='中文 new';link.address='https://other.invalid/new';link.fragment='detail'
    internal=paragraph.add_hyperlink('Jump',fragment='Bookmark')
    output=BytesIO();document.save(output);output.seek(0);restored=Document(output).paragraphs[0].hyperlinks
    assert restored[0].text=='中文 new' and restored[0].url=='https://other.invalid/new#detail' and restored[0].runs[0].bold
    assert restored[1].address=='' and restored[1].fragment=='Bookmark'

def test_shared_relationship_retained_until_last_reference_removed():
    "Verify shared relationship retained until last reference removed."
    document=Document();paragraph=document.add_paragraph()
    a=paragraph.add_hyperlink('a','https://same.invalid');b=paragraph.add_hyperlink('b','https://same.invalid')
    shared=a._hyperlink.rId;assert shared==b._hyperlink.rId
    a.address='https://new.invalid';assert b.address=='https://same.invalid' and shared in document.part.rels
    b.remove();assert shared not in document.part.rels and len(paragraph.hyperlinks)==1
    a.remove();a.remove();assert paragraph.hyperlinks==[]


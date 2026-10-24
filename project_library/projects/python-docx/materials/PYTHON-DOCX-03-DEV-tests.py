"Public regression tests for PYTHON-DOCX-03-DEV; the complete business contract is in materials/PYTHON-DOCX-03-DEV-contract.md."
from io import BytesIO
from zipfile import ZipFile
import unittest
from docx import Document

def test_foreign_empty_reversed_and_header_ranges_have_no_side_effects():
    "Verify foreign empty reversed and header ranges have no side effects."
    target=Document();a=target.add_paragraph('a').runs[0];b=target.add_paragraph('b').runs[0]
    foreign=Document();other=foreign.add_paragraph('foreign').runs[0]
    header=target.sections[0].header.paragraphs[0].add_run('header')
    before=target._element.xml;foreign_before=foreign._element.xml
    for runs in ([other],[],[b,a],[a,other],[header],[a,a]):
        with unittest.TestCase().assertRaises(ValueError):target.add_comment(runs,'invalid')
        assert target._element.xml==before and foreign._element.xml==foreign_before
    output=BytesIO();target.save(output)
    with ZipFile(output) as archive:assert 'word/comments.xml' not in archive.namelist()

def test_same_document_cross_paragraph_comment_saves_with_consistent_ids():
    "Verify same document cross paragraph comment saves with consistent ids."
    document=Document();first=document.add_paragraph('first').runs[0];last=document.add_paragraph('last').runs[0]
    comment=document.add_comment([first,last],'review',author='Reviewer')
    assert comment.text=='review'
    output=BytesIO();document.save(output);output.seek(0);restored=Document(output)
    assert len(restored.comments)==1 and next(iter(restored.comments)).text=='review'


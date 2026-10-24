"""独立业务验收，只读取当前候选，不联网或调用模型。"""
import json
import os
from pathlib import Path
root = Path(os.environ['PROJECT_CANDIDATE_ROOT']).resolve()
from contextlib import ExitStack
from importlib.util import module_from_spec, spec_from_file_location
spec = spec_from_file_location('frozen_native_control', Path(__file__).with_name('native_regression.py'))
control = module_from_spec(spec)
spec.loader.exec_module(control)
cleanup = ExitStack()
suite = control.SUITES['python-docx']
control._bind_source(root / 'library', Path(suite.get('import_root', '.')), cleanup, distribution=suite.get('distribution'))
import sys
sys.path.insert(0, str(root / 'library' / suite.get('import_root', '.')))
import docx
from io import BytesIO
from docx import Document
business=True
document=Document();row=document.add_table(rows=1,cols=1).rows[0]
row.repeat_header=True;row.allow_split_across_pages=False
output=BytesIO();document.save(output);output.seek(0);row=Document(output).tables[0].rows[0]
business &= row.repeat_header is True and row.allow_split_across_pages is False
row.allow_split_across_pages=None;business &= row.allow_split_across_pages is None and row.repeat_header is True

print(json.dumps({'candidate_import': Path(docx.__file__).resolve().is_relative_to(root / 'library'), 'business': bool(business)}))

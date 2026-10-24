"""独立业务验收，只读取当前候选，不联网或调用模型。"""
import json
import os
from pathlib import Path
root = Path(os.environ['PROJECT_CANDIDATE_ROOT']).resolve()
from contextlib import ExitStack
from importlib.util import module_from_spec, spec_from_file_location
spec = spec_from_file_location('frozen_native_control', Path(__file__).with_name('native_regression.py'))
control = module_from_spec(spec)
spec.loader.exec_module(control)
cleanup = ExitStack()
suite = control.SUITES['python-docx']
control._bind_source(root / 'library', Path(suite.get('import_root', '.')), cleanup, distribution=suite.get('distribution'))
import sys
sys.path.insert(0, str(root / 'library' / suite.get('import_root', '.')))
import docx
from io import BytesIO
from docx import Document
business=True
document=Document();paragraph=document.add_paragraph()
link=paragraph.add_hyperlink('email','mailto:test@example.invalid');old=link._hyperlink.rId
link.fragment='internal';link.address=''
business &= link.address=='' and link.fragment=='internal' and old not in document.part.rels
output=BytesIO();document.save(output);output.seek(0)
business &= Document(output).paragraphs[0].hyperlinks[0].text=='email'

print(json.dumps({'candidate_import': Path(docx.__file__).resolve().is_relative_to(root / 'library'), 'business': bool(business)}))

"Public regression tests for PYTHON-PPTX-03-DEV; the complete business contract is in materials/PYTHON-PPTX-03-DEV-contract.md."
from io import BytesIO
from pptx import Presentation

def test_textbox_and_table_accessibility_metadata_roundtrip():
    "Verify textbox and table accessibility metadata roundtrip."
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes=[slide.shapes.add_textbox(1,2,300000,400000),slide.shapes.add_table(1,1,5,6,500000,500000)]
    before=[(shape.shape_id,shape.left,shape.width) for shape in shapes]
    for index,shape in enumerate(shapes):shape.alternative_text='说明 & '+str(index);shape.title='Figure '+str(index)
    output=BytesIO();presentation.save(output);output.seek(0);restored=Presentation(output).slides[0].shapes
    assert [(shape.shape_id,shape.left,shape.width) for shape in restored]==before
    assert [shape.alternative_text for shape in restored]==['说明 & 0','说明 & 1']
    assert restored[1].title=='Figure 1'

def test_none_vs_empty_and_other_nonvisual_properties_stay():
    "Verify none vs empty and other nonvisual properties stay."
    presentation=Presentation();shape=presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(1,2,30,40)
    name=shape.name;assert shape.alternative_text is None
    shape.alternative_text='';shape.title='keep';assert shape.alternative_text==''
    shape.alternative_text=None
    assert shape.alternative_text is None and shape.title=='keep' and shape.name==name


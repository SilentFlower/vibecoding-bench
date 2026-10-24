"Public regression tests for PYTHON-PPTX-06-DEV; the complete business contract is in materials/PYTHON-PPTX-06-DEV-contract.md."
from io import BytesIO
import unittest
from pptx import Presentation
from pptx.layout import align_shapes,distribute_shapes

def test_align_and_equal_gaps_with_unequal_widths_roundtrip():
    "Verify align and equal gaps with unequal widths roundtrip."
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes=[slide.shapes.add_textbox(left,top,width,100000) for left,top,width in ((0,200000,100000),(270000,80000,200000),(700000,400000,100000))]
    ids=[shape.shape_id for shape in shapes];align_shapes(shapes,'top');distribute_shapes(shapes,'x')
    assert [shape.left for shape in shapes]==[0,300000,700000] and all(shape.top==80000 for shape in shapes)
    output=BytesIO();presentation.save(output);output.seek(0);restored=Presentation(output).slides[0].shapes
    assert [shape.left for shape in restored]==[0,300000,700000] and [shape.shape_id for shape in restored]==ids

def test_centers_empty_and_invalid_rotated_or_mixed_parent_atomic():
    "Verify centers empty and invalid rotated or mixed parent atomic."
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    a=slide.shapes.add_textbox(0,0,100,200);b=slide.shapes.add_textbox(500,500,200,100)
    align_shapes([a,b],'center_x');assert a.left*2+a.width==b.left*2+b.width
    assert align_shapes([],'left')==()
    b.rotation=10;before=(a.left,a.top,b.left,b.top)
    with unittest.TestCase().assertRaises(ValueError):align_shapes([a,b],'left')
    assert (a.left,a.top,b.left,b.top)==before


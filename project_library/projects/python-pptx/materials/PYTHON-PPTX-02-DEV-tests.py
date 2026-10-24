"Public regression tests for PYTHON-PPTX-02-DEV; the complete business contract is in materials/PYTHON-PPTX-02-DEV-contract.md."
from io import BytesIO
from zipfile import ZipFile
import unittest
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

def test_fade_timed_advance_and_push_roundtrip():
    "Verify fade timed advance and push roundtrip."
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    assert slide.transition.configuration is None
    slide.transition.configure('fade',advance_after_ms=3000,advance_on_click=False,speed='slow')
    output=BytesIO();presentation.save(output)
    with ZipFile(output) as archive:assert b'advTm="3000"' in archive.read('ppt/slides/slide1.xml')
    output.seek(0);slide=Presentation(output).slides[0]
    assert slide.transition.configuration=={'effect':'fade','direction':None,'speed':'slow','advance_on_click':False,'advance_after_ms':3000}
    slide.transition.configure('push',direction='u')
    assert slide.transition.configuration['direction']=='u'

def test_invalid_or_unknown_changes_are_atomic_and_clear_is_explicit():
    "Verify invalid or unknown changes are atomic and clear is explicit."
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.transition.configure('wipe',direction='r');before=slide._element.xml
    with unittest.TestCase().assertRaises(ValueError):slide.transition.configure('fade',direction='l')
    assert slide._element.xml==before
    slide.transition.clear();node=OxmlElement('p:transition');node.append(OxmlElement('p:wheel'));slide._element.append(node)
    before=slide._element.xml
    with unittest.TestCase().assertRaises(ValueError):slide.transition.configure('fade')
    assert slide._element.xml==before;slide.transition.clear();assert slide.transition.configuration is None


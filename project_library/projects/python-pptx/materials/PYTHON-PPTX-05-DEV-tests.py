"""文字流向是bodyPr方向，不能用旋转文本框冒充竖排。"""
from io import BytesIO
import unittest
from pptx import Presentation

def test_vertical_east_asian_and_270_roundtrip_leave_shape_rotation():
    """正文/内边距/形状旋转保持，只更改文字流向。@return None。"""
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    for index,value in enumerate(('vert','eaVert','vert270')):
        shape=slide.shapes.add_textbox(index*400000,0,300000,400000);shape.text='中文 direction'
        shape.text_frame.text_direction=value
    output=BytesIO();presentation.save(output);output.seek(0);shapes=Presentation(output).slides[0].shapes
    assert [shape.text_frame.text_direction for shape in shapes]==['vert','eaVert','vert270']
    assert all(shape.rotation==0 and shape.text=='中文 direction' for shape in shapes)

def test_none_inheritance_horizontal_and_invalid_atomic():
    """None移除直接属性，不覆盖其他bodyPr。@return None。"""
    presentation=Presentation();frame=presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(0,0,300000,300000).text_frame
    margin=frame.margin_left;assert frame.text_direction is None
    frame.text_direction='horz';before=frame._bodyPr.xml
    with unittest.TestCase().assertRaises(ValueError):frame.text_direction='rotated'
    assert frame._bodyPr.xml==before
    frame.text_direction=None;assert frame.text_direction is None and frame.margin_left==margin


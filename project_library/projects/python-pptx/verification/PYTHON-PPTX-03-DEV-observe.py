"""独立业务验收，只读取当前候选，不联网或调用模型。"""
import json
import os
from pathlib import Path
root = Path(os.environ['PROJECT_CANDIDATE_ROOT']).resolve()
from contextlib import ExitStack
from importlib.util import module_from_spec, spec_from_file_location
spec = spec_from_file_location('frozen_native_control', Path(__file__).with_name('native_regression.py'))
control = module_from_spec(spec)
spec.loader.exec_module(control)
cleanup = ExitStack()
suite = control.SUITES['python-pptx']
control._bind_source(root / 'library', Path(suite.get('import_root', '.')), cleanup, distribution=suite.get('distribution'))
import sys
sys.path.insert(0, str(root / 'library' / suite.get('import_root', '.')))
import pptx
from io import BytesIO
from pptx import Presentation
business=True
presentation=Presentation();shape=presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(0,0,400000,300000)
shape.alternative_text='diagram';shape.title='Business'
output=BytesIO();presentation.save(output);output.seek(0);again=Presentation(output).slides[0].shapes[0]
business &= again.alternative_text=='diagram' and again.title=='Business'

print(json.dumps({'candidate_import': Path(pptx.__file__).resolve().is_relative_to(root / 'library'), 'business': bool(business)}))

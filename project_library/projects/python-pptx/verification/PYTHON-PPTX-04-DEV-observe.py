"""独立业务验收，只读取当前候选，不联网或调用模型。"""
import json
import os
from pathlib import Path
root = Path(os.environ['PROJECT_CANDIDATE_ROOT']).resolve()
from contextlib import ExitStack
from importlib.util import module_from_spec, spec_from_file_location
spec = spec_from_file_location('frozen_native_control', Path(__file__).with_name('native_regression.py'))
control = module_from_spec(spec)
spec.loader.exec_module(control)
cleanup = ExitStack()
suite = control.SUITES['python-pptx']
control._bind_source(root / 'library', Path(suite.get('import_root', '.')), cleanup, distribution=suite.get('distribution'))
import sys
sys.path.insert(0, str(root / 'library' / suite.get('import_root', '.')))
import pptx
from pptx import Presentation
business=True
presentation=Presentation();paragraph=presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(0,0,300000,300000).text_frame.paragraphs[0]
paragraph.set_bullet('number',scheme='alphaLcParenR',start=3)
business &= paragraph.bullet=={'kind':'number','scheme':'alphaLcParenR','start':3}
paragraph.set_bullet('character',character='-')
business &= paragraph.bullet=={'kind':'character','character':'-'} and len(paragraph._p.pPr)==1

print(json.dumps({'candidate_import': Path(pptx.__file__).resolve().is_relative_to(root / 'library'), 'business': bool(business)}))
